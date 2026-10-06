"""Build data/campgrounds.js from the three Excel files in source/.
Usage: pip install -r requirements.txt && python tools/build_data.py
Locations: optional Latitude/Longitude columns win; otherwise the ZIP-code centre
(precision "zip"), then the city (precision "city"). Parks that cannot be placed are listed, not drawn. US only."""
import json, re, datetime
from collections import Counter, defaultdict
from pathlib import Path
import zipcodes, geonamescache
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
US = {'AL':'Alabama','AK':'Alaska','AZ':'Arizona','AR':'Arkansas','CA':'California','CO':'Colorado','CT':'Connecticut','DE':'Delaware','DC':'District of Columbia','FL':'Florida','GA':'Georgia','HI':'Hawaii','ID':'Idaho','IL':'Illinois','IN':'Indiana','IA':'Iowa','KS':'Kansas','KY':'Kentucky','LA':'Louisiana','ME':'Maine','MD':'Maryland','MA':'Massachusetts','MI':'Michigan','MN':'Minnesota','MS':'Mississippi','MO':'Missouri','MT':'Montana','NE':'Nebraska','NV':'Nevada','NH':'New Hampshire','NJ':'New Jersey','NM':'New Mexico','NY':'New York','NC':'North Carolina','ND':'North Dakota','OH':'Ohio','OK':'Oklahoma','OR':'Oregon','PA':'Pennsylvania','RI':'Rhode Island','SC':'South Carolina','SD':'South Dakota','TN':'Tennessee','TX':'Texas','UT':'Utah','VT':'Vermont','VA':'Virginia','WA':'Washington','WV':'West Virginia','WI':'Wisconsin','WY':'Wyoming'}
CA = {'BC':'British Columbia','AB':'Alberta','ON':'Ontario','QC':'Quebec','NB':'New Brunswick','NS':'Nova Scotia','PE':'Prince Edward Island','YT':'Yukon Territory'}
ABBR = {**US, **CA}
CA_ADMIN = {'01':'Alberta','02':'British Columbia','04':'New Brunswick','07':'Nova Scotia','08':'Ontario','09':'Prince Edward Island','10':'Quebec','12':'Yukon Territory'}
FULL2ABBR = {v: k for k, v in US.items()}
NETS = [('tt', 'thousand_trails.xlsx'), ('en', 'encore.xlsx'), ('pa', 'passport_america.xlsx')]
gc = geonamescache.GeonamesCache()
cities = defaultdict(list)
for c in gc.get_cities().values():
    cities[c['name'].lower()].append(c)

def find_col(header, *names):
    for i, h in enumerate(header):
        if h and str(h).strip().lower() in names: return i

def load(path):
    ws = load_workbook(path, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r and any(str(x).strip().lower() == 'name' for x in r if x))
    head = rows[h]
    col = {k: find_col(head, *v) for k, v in dict(name=('name',), state=('state',), addr=('address',), city=('city or town', 'city / town'), zip=('zipcode', 'zip'), notes=('notes',), lat=('latitude', 'lat'), lng=('longitude', 'lng', 'lon')).items()}
    for r in rows[h + 1:]:
        if not r[col['name']] or (col['state'] is not None and not r[col['state']]): continue
        get = lambda k: (str(r[col[k]]).strip() if col[k] is not None and r[col[k]] not in (None, '') else '')
        yield {k: get(k) for k in col}

def geocode(rec, state_full):
    if rec['lat'] and rec['lng']: return float(rec['lat']), float(rec['lng']), 'exact'
    z = re.match(r'\d{5}', rec['zip'])
    if z and state_full in FULL2ABBR:
        for m in zipcodes.matching(z.group()):
            if m['state'] == FULL2ABBR[state_full] and float(m['lat'] or 0) and float(m['long'] or 0): return float(m['lat']), float(m['long']), 'zip'
    cc = rec['city'].split('/')[0].strip().lower()
    for c in sorted(cities.get(cc, []), key=lambda c: -c['population']):
        if state_full in FULL2ABBR and c['countrycode'] == 'US' and c['admin1code'] == FULL2ABBR[state_full]: return c['latitude'], c['longitude'], 'city'
        if state_full not in FULL2ABBR and c['countrycode'] in ('CA', 'MX') and (c['countrycode'] == 'MX' or CA_ADMIN.get(c['admin1code']) == state_full): return c['latitude'], c['longitude'], 'city'
    return None

def main():
    out = {k: [] for k, _ in NETS}; skipped = []; bad = []
    for net, fn in NETS:
        seen = set()
        for rec in load(ROOT / 'source' / fn):
            if rec['notes'].lower().startswith('verify'):
                skipped.append((net, rec['name'], rec['notes'])); continue
            st = ABBR.get(rec['state'].upper(), rec['state'])
            if st not in US.values():
                skipped.append((net, rec['name'], 'not in the United States')); continue
            key = (rec['name'].lower(), rec['zip'][:5])
            if key in seen or 'annex' in rec['name'].lower():
                skipped.append((net, rec['name'], 'duplicate or annex of another listing')); continue
            seen.add(key)
            g = geocode(rec, st)
            out[net].append(dict(n=rec['name'], s=st, c=rec['city'], a=rec['addr'], z=rec['zip'], g=g))
    for net, v in out.items():
        bad += [(net, r['n'], r['s']) for r in v if not r['g']]
        out[net] = [dict(n=r['n'], s=r['s'], c=r['c'], a=r['a'], z=r['z'], lat=round(r['g'][0], 5), lng=round(r['g'][1], 5), p=r['g'][2]) for r in v if r['g']]
    # state bounds: every US ZIP centre, plus the campground points for other regions
    b = {}
    for z in zipcodes.list_all():
        s = US.get(z['state'])
        la, lo = float(z['lat'] or 0), float(z['long'] or 0)
        if s and la and lo and z['zip_code_type'] != 'MILITARY':
            if s == 'Alaska' and lo > 0: continue
            x = b.setdefault(s, [la, lo, la, lo]); x[:] = [min(x[0], la), min(x[1], lo), max(x[2], la), max(x[3], lo)]
    for v in out.values():
        for r in v:
            if r['s'] not in US.values():
                x = b.setdefault(r['s'], [r['lat'], r['lng'], r['lat'], r['lng']]); x[:] = [min(x[0], r['lat']), min(x[1], r['lng']), max(x[2], r['lat']), max(x[3], r['lng'])]
    data = dict(built=datetime.date.today().isoformat(), networks=out, bounds={k: [round(x, 3) for x in v] for k, v in sorted(b.items())}, abbr={v: k for k, v in US.items()})
    (ROOT / 'data').mkdir(exist_ok=True)
    (ROOT / 'data' / 'campgrounds.js').write_text('window.CAMPGROUNDS = ' + json.dumps(data, separators=(',', ':'), ensure_ascii=False) + ';\n', encoding='utf-8')
    for net, v in out.items(): print(net, len(v), dict(Counter(r['p'] for r in v)))
    print('skipped:', [x for x in skipped if x[2] != 'not in the United States'], '| non-US skipped:', sum(x[2] == 'not in the United States' for x in skipped)); print('could not place:', bad)
    print('states:', len(data['bounds']))

main()
